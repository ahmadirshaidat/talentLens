"""Shared pytest fixtures. All CV content is fake."""

import io
from collections.abc import Callable

import docx
import pymupdf
import pytest

from app.storage.vector_store import VectorStore


@pytest.fixture
def make_docx() -> Callable[..., bytes]:
    """Build a DOCX in memory: header line, paragraphs, and an optional table."""

    def _make(
        paragraphs: list[str],
        table: list[list[str]] | None = None,
        header: str | None = None,
    ) -> bytes:
        document = docx.Document()
        if header:
            document.sections[0].header.paragraphs[0].text = header
        for text in paragraphs:
            document.add_paragraph(text)
        if table:
            t = document.add_table(rows=len(table), cols=len(table[0]))
            for r, row in enumerate(table):
                for c, value in enumerate(row):
                    t.cell(r, c).text = value
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    return _make


@pytest.fixture
def make_pdf() -> Callable[[list[str]], bytes]:
    """Build a PDF in memory with one page per string (Latin text only — base-14 font)."""

    def _make(pages: list[str]) -> bytes:
        doc = pymupdf.open()
        for text in pages:
            page = doc.new_page()
            if text:
                page.insert_text((72, 72), text, fontsize=11)
        data = doc.tobytes()
        doc.close()
        return data

    return _make


@pytest.fixture
def vector_store(tmp_path) -> VectorStore:
    return VectorStore(tmp_path / "chroma")


class FakeEncoder:
    """Stands in for SentenceTransformer: deterministic 4-dim vectors from keywords."""

    KEYWORDS = ("python", "java", "sales", "عربي")

    def __init__(self) -> None:
        self.calls: list[dict] = []

    def encode(self, sentences: list[str], **kwargs):
        self.calls.append(kwargs)
        vectors = []
        for s in sentences:
            v = [1.0 if k in s.lower() else 0.0 for k in self.KEYWORDS]
            if not any(v):
                v = [0.25, 0.25, 0.25, 0.25]
            norm = sum(x * x for x in v) ** 0.5
            vectors.append([x / norm for x in v])
        return vectors


@pytest.fixture
def fake_encoder() -> FakeEncoder:
    return FakeEncoder()
