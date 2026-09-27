"""Generate fake CVs (fake names only, Arabic + English) and a labeled eval dataset.

Writes:
    sample_cvs/cv_XX.docx | cv_XX.pdf      — 16 fake CVs (English PDF/DOCX, Arabic DOCX)
    eval/datasets/sample_queries.json      — recruiter queries with graded relevant CVs

Usage (from ai-service/):
    python scripts/generate_sample_cvs.py [--out sample_cvs]

All people, companies, emails and phone numbers are invented.
"""

import argparse
import io
import json
import sys
import textwrap
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import docx
import pymupdf
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
THIS_YEAR = date.today().year


@dataclass
class Job:
    title: str
    company: str
    years: int  # duration; jobs are laid out back-to-back ending "Present"
    bullets: list[str]


@dataclass
class Person:
    cid: str
    lang: str  # "en" | "ar"
    fmt: str  # "docx" | "pdf"
    name: str
    email: str
    phone: str
    location: str
    summary: str
    jobs: list[Job]
    skills: list[str]
    education: list[str]
    languages: list[str]
    projects: list[str] = field(default_factory=list)
    skills_as_table: bool = False

    @property
    def total_years(self) -> int:
        return sum(j.years for j in self.jobs)


PEOPLE: list[Person] = [
    Person(
        "cv_01", "en", "docx", "Lina Haddad", "lina.haddad@example.com", "+962 79 000 1101",
        "Amman, Jordan",
        "Backend developer with 6 years of experience building Python web services and APIs.",
        [
            Job("Senior Backend Developer", "Nimbus Apps", 4, [
                "Designed REST APIs in Python and Django serving 2M requests per day",
                "Moved services to Docker and cut deployment time by 60%",
                "Mentored 3 junior developers on code reviews and testing",
            ]),
            Job("Backend Developer", "Cedar Soft", 2, [
                "Built PostgreSQL data models and reporting endpoints",
                "Wrote unit tests with pytest reaching 85% coverage",
            ]),
        ],
        ["Python", "Django", "PostgreSQL", "Docker", "REST APIs", "Git"],
        ["B.Sc. Computer Science, Example University, 2019"],
        ["Arabic (native)", "English (fluent)"],
        ["Open-source Django package for audit logging"],
        skills_as_table=True,
    ),
    Person(
        "cv_02", "en", "pdf", "Omar Nasser", "omar.nasser@example.com", "+962 78 000 1202",
        "Irbid, Jordan",
        "Senior .NET developer with 8 years of experience in C#, ASP.NET Core and SQL Server.",
        [
            Job("Senior .NET Developer", "BlueFalcon Systems", 5, [
                "Led a team of 5 engineers building an ASP.NET Core banking portal",
                "Optimized SQL Server queries and reduced report time from 40s to 3s",
                "Deployed microservices to Azure App Service",
            ]),
            Job(".NET Developer", "Olive Tech", 3, [
                "Maintained C# desktop and web applications",
                "Integrated payment gateways through REST APIs",
            ]),
        ],
        ["C#", "ASP.NET Core", "SQL Server", "Azure", "Microservices", "Git"],
        ["B.Sc. Software Engineering, Example University, 2017"],
        ["Arabic", "English"],
    ),
    Person(
        "cv_03", "ar", "docx", "سارة خليل", "sara.khalil@example.com", "+962 77 000 1303",
        "عمّان، الأردن",
        "مطورة واجهات أمامية لديها خبرة 3 سنوات في بناء تطبيقات ويب باستخدام React و TypeScript.",
        [
            Job("مطورة واجهات أمامية", "شركة نجمة الرقمية", 2, [
                "تطوير واجهات مستخدم تفاعلية باستخدام React و TypeScript",
                "تحسين سرعة تحميل الصفحات بنسبة 40%",
            ]),
            Job("مطورة ويب مبتدئة", "شركة الأفق للبرمجيات", 1, [
                "بناء صفحات ويب متجاوبة باستخدام JavaScript و CSS",
            ]),
        ],
        ["React", "JavaScript", "TypeScript", "Git", "Figma"],
        ["بكالوريوس علم الحاسوب، جامعة المثال، 2022"],
        ["العربية", "الإنجليزية"],
        ["متجر إلكتروني باستخدام React"],
    ),
    Person(
        "cv_04", "ar", "docx", "يوسف عوض", "yousef.awad@example.com", "+962 79 000 1404",
        "الزرقاء، الأردن",
        "محاسب قانوني بخبرة 7 سنوات في إعداد القوائم المالية والتدقيق.",
        [
            Job("محاسب أول", "مجموعة الواحة التجارية", 4, [
                "إعداد القوائم المالية الشهرية والسنوية",
                "إدارة الحسابات الدائنة والمدينة باستخدام Excel",
            ]),
            Job("محاسب", "شركة الريادة للتجارة", 3, [
                "مطابقة الحسابات البنكية وإعداد التقارير الضريبية",
            ]),
        ],
        ["المحاسبة", "Excel", "التقارير المالية", "التدقيق"],
        ["بكالوريوس محاسبة، جامعة المثال، 2018"],
        ["العربية"],
    ),
    Person(
        "cv_05", "en", "pdf", "Hala Mansour", "hala.mansour@example.com", "+962 78 000 1505",
        "Amman, Jordan",
        "Data scientist with 4 years of experience in machine learning and NLP.",
        [
            Job("Data Scientist", "Insight Labs", 3, [
                "Trained machine learning models with Python, Pandas and TensorFlow",
                "Built an Arabic NLP classifier for customer support tickets",
            ]),
            Job("Data Analyst", "Cedar Soft", 1, [
                "Wrote SQL queries and dashboards for the sales team",
            ]),
        ],
        ["Python", "Machine Learning", "NLP", "Pandas", "TensorFlow", "SQL"],
        ["M.Sc. Data Science, Example University, 2021"],
        ["English", "French", "Arabic"],
    ),
    Person(
        "cv_06", "en", "docx", "Kareem Saleh", "kareem.saleh@example.com", "+962 79 000 1606",
        "Aqaba, Jordan",
        "Sales manager with 9 years of experience growing B2B accounts.",
        [
            Job("Sales Manager", "Desert Trade Co.", 5, [
                "Managed a sales team of 8 and grew revenue by 35% in two years",
                "Negotiated contracts with 40+ corporate clients",
            ]),
            Job("Sales Representative", "Harbor Supplies", 4, [
                "Handled customer service escalations and key accounts",
                "Prepared weekly sales reports in Excel",
            ]),
        ],
        ["Sales", "Customer Service", "Communication", "Leadership", "Excel"],
        ["B.A. Business Administration, Example University, 2016"],
        ["Arabic", "English"],
    ),
    Person(
        "cv_07", "ar", "docx", "رنا عودة", "rana.odeh@example.com", "+962 77 000 1707",
        "إربد، الأردن",
        "مهندسة بيانات بخبرة 5 سنوات في تحليل البيانات وبناء لوحات Power BI.",
        [
            Job("مهندسة بيانات", "شركة البيان للحلول", 3, [
                "بناء خطوط معالجة البيانات باستخدام Python و SQL",
                "تصميم لوحات تقارير تفاعلية في Power BI للإدارة العليا",
            ]),
            Job("محللة بيانات", "شركة الأفق للبرمجيات", 2, [
                "تحليل البيانات وإعداد تقارير أداء المبيعات",
            ]),
        ],
        ["Python", "SQL", "Power BI", "تحليل البيانات", "Excel"],
        ["بكالوريوس نظم المعلومات، جامعة المثال، 2020"],
        ["العربية", "الإنجليزية"],
    ),
    Person(
        "cv_08", "en", "pdf", "Tariq Hamdan", "tariq.hamdan@example.com", "+962 78 000 1808",
        "Amman, Jordan",
        "DevOps engineer with 6 years of experience automating cloud infrastructure.",
        [
            Job("DevOps Engineer", "Skyline Cloud", 4, [
                "Ran production Kubernetes clusters on AWS for 30 microservices",
                "Wrote Python automation scripts and CI/CD pipelines",
            ]),
            Job("Linux System Administrator", "Olive Tech", 2, [
                "Managed Linux servers and Docker based deployments",
            ]),
        ],
        ["Docker", "Kubernetes", "AWS", "Linux", "Python", "Git"],
        ["B.Sc. Computer Engineering, Example University, 2019"],
        ["Arabic", "English"],
    ),
    Person(
        "cv_09", "en", "docx", "Maya Barakat", "maya.barakat@example.com", "+962 79 000 1909",
        "Salt, Jordan",
        "Junior Python developer with 1 year of experience in Flask web apps.",
        [
            Job("Junior Python Developer", "Tiny Start", 1, [
                "Built small Flask APIs and fixed bugs in the admin dashboard",
                "Used Git and GitHub for pull requests",
            ]),
        ],
        ["Python", "Flask", "Git", "SQL"],
        ["B.Sc. Computer Science, Example University, 2025"],
        ["English", "Arabic"],
        ["Graduation project: course registration system in Flask"],
    ),
    Person(
        "cv_10", "ar", "docx", "فيصل قاسم", "faisal.qasem@example.com", "+962 77 000 2010",
        "عمّان، الأردن",
        "مطور .NET بخبرة 4 سنوات في تطوير تطبيقات الويب باستخدام C# و ASP.NET Core.",
        [
            Job("مطور .NET", "شركة الصقر للأنظمة", 3, [
                "تطوير أنظمة إدارة داخلية باستخدام ASP.NET Core و SQL Server",
                "كتابة واجهات REST APIs وربطها مع تطبيقات الجوال",
            ]),
            Job("متدرب برمجيات", "شركة نجمة الرقمية", 1, [
                "المشاركة في تطوير وحدات C# واختبارها",
            ]),
        ],
        ["C#", "ASP.NET Core", "SQL Server", "REST APIs", "Git"],
        ["بكالوريوس هندسة البرمجيات، جامعة المثال، 2021"],
        ["العربية", "الإنجليزية"],
    ),
    Person(
        "cv_11", "en", "pdf", "Nour Jaber", "nour.jaber@example.com", "+962 78 000 2111",
        "Amman, Jordan",
        "Digital marketing specialist with 5 years of experience in SEO and content.",
        [
            Job("Digital Marketing Specialist", "BrightAds", 3, [
                "Grew organic traffic 3x through SEO and content strategy",
                "Ran paid campaigns in Arabic, English and French markets",
            ]),
            Job("Marketing Coordinator", "Harbor Supplies", 2, [
                "Managed social media channels and newsletters",
            ]),
        ],
        ["Marketing", "SEO", "Communication", "Excel"],
        ["B.A. Marketing, Example University, 2020"],
        ["Arabic", "English", "French"],
    ),
    Person(
        "cv_12", "en", "docx", "Ziad Fares", "ziad.fares@example.com", "+962 79 000 2212",
        "Madaba, Jordan",
        "Mobile developer with 3 years of experience building Flutter and Android apps.",
        [
            Job("Mobile Developer", "AppNest", 3, [
                "Shipped 4 Flutter apps to the App Store and Google Play",
                "Wrote native Android modules in Kotlin",
            ]),
        ],
        ["Flutter", "Android", "Kotlin", "Git", "REST APIs"],
        ["B.Sc. Computer Science, Example University, 2022"],
        ["Arabic", "English"],
        ["Food delivery app in Flutter with 10k downloads"],
    ),
    Person(
        "cv_13", "ar", "docx", "دانة الشامي", "dana.shami@example.com", "+962 77 000 2313",
        "عمّان، الأردن",
        "أخصائية توظيف بخبرة 6 سنوات في استقطاب الكفاءات وإجراء المقابلات.",
        [
            Job("أخصائية توظيف", "مجموعة الواحة التجارية", 4, [
                "إدارة عملية التوظيف الكاملة لأكثر من 120 وظيفة سنوياً",
                "إجراء المقابلات وتقييم المرشحين",
            ]),
            Job("منسقة موارد بشرية", "شركة الريادة للتجارة", 2, [
                "تنظيم ملفات الموظفين وبرامج التهيئة",
            ]),
        ],
        ["التوظيف", "مهارات التواصل", "Excel"],
        ["بكالوريوس إدارة الموارد البشرية، جامعة المثال، 2019"],
        ["العربية", "الإنجليزية"],
    ),
    Person(
        "cv_14", "en", "pdf", "Samer Zoubi", "samer.zoubi@example.com", "+962 78 000 2414",
        "Irbid, Jordan",
        "Java developer with 7 years of experience in Spring Boot microservices.",
        [
            Job("Senior Java Developer", "BlueFalcon Systems", 4, [
                "Built Spring Boot microservices with MySQL and Redis",
                "Introduced contract testing between services",
            ]),
            Job("Java Developer", "Cedar Soft", 3, [
                "Maintained Java web applications and REST APIs",
            ]),
        ],
        ["Java", "Spring", "Microservices", "MySQL", "Redis", "Docker"],
        ["B.Sc. Computer Science, Example University, 2018"],
        ["Arabic", "English"],
    ),
    Person(
        "cv_15", "en", "docx", "Leen Abbas", "leen.abbas@example.com", "+962 79 000 2515",
        "Amman, Jordan",
        "UI/UX designer with 4 years of experience designing web and mobile products.",
        [
            Job("UI/UX Designer", "AppNest", 3, [
                "Designed mobile app flows and design systems in Figma",
                "Ran usability tests with 50+ users",
            ]),
            Job("Graphic Designer", "BrightAds", 1, [
                "Created campaign visuals in Photoshop",
            ]),
        ],
        ["Figma", "Photoshop", "Communication"],
        ["B.A. Graphic Design, Example University, 2021"],
        ["Arabic", "English"],
    ),
    Person(
        "cv_16", "ar", "docx", "بسام الخوري", "bassam.khoury@example.com", "+962 77 000 2616",
        "عمّان، الأردن",
        "مدير مشاريع تقنية بخبرة 10 سنوات في قيادة فرق Agile وتسليم المشاريع في وقتها.",
        [
            Job("مدير مشاريع", "شركة الصقر للأنظمة", 6, [
                "قيادة 4 فرق تطوير باستخدام منهجية Agile و Scrum",
                "إدارة ميزانيات مشاريع تتجاوز مليون دينار",
            ]),
            Job("منسق مشاريع", "شركة البيان للحلول", 4, [
                "متابعة الجداول الزمنية والتواصل مع العملاء",
            ]),
        ],
        ["إدارة المشاريع", "Agile", "القيادة", "مهارات التواصل"],
        ["ماجستير إدارة الأعمال، جامعة المثال، 2015"],
        ["العربية", "الإنجليزية"],
    ),
]

# Graded relevance: 2 = strong match, 1 = partial match.
QUERIES: list[dict] = [
    {"query": "Python developer with 5+ years of experience",
     "relevant": {"cv_01": 2, "cv_07": 1, "cv_08": 1}},
    {"query": ".NET developer C# SQL Server", "relevant": {"cv_02": 2, "cv_10": 2}},
    {"query": "مطور .NET خبرة 5 سنوات", "relevant": {"cv_02": 2}},
    {"query": "frontend React developer", "relevant": {"cv_03": 2}},
    {"query": "مطور واجهات أمامية React", "relevant": {"cv_03": 2}},
    {"query": "machine learning data scientist", "relevant": {"cv_05": 2, "cv_07": 1}},
    {"query": "محاسب خبرة في Excel", "relevant": {"cv_04": 2}},
    {"query": "sales manager with strong communication skills who speaks English",
     "relevant": {"cv_06": 2, "cv_11": 1}},
    {"query": "DevOps engineer Kubernetes AWS", "relevant": {"cv_08": 2}},
    {"query": "Java Spring microservices", "relevant": {"cv_14": 2}},
    {"query": "أخصائي توظيف موارد بشرية", "relevant": {"cv_13": 2}},
    {"query": "project manager agile scrum", "relevant": {"cv_16": 2}},
    {"query": "mobile app developer Flutter", "relevant": {"cv_12": 2}},
    {"query": "UI UX designer Figma", "relevant": {"cv_15": 2, "cv_03": 1}},
    {"query": "digital marketing SEO French", "relevant": {"cv_11": 2}},
    {"query": "data analyst Power BI SQL", "relevant": {"cv_07": 2, "cv_05": 1}},
]

HEADINGS = {
    "en": {
        "summary": "Professional Summary", "experience": "Work Experience",
        "skills": "Skills", "education": "Education", "languages": "Languages",
        "projects": "Projects", "present": "Present", "location": "Location",
    },
    "ar": {
        "summary": "الملخص المهني", "experience": "الخبرات العملية",
        "skills": "المهارات", "education": "التعليم", "languages": "اللغات",
        "projects": "المشاريع", "present": "حتى الآن", "location": "العنوان",
    },
}


def _job_lines(p: Person) -> list[tuple[str, list[str]]]:
    """(heading line, bullets) per job, most recent first, dates ending at Present."""
    h = HEADINGS[p.lang]
    end: int | str = h["present"]
    end_year = THIS_YEAR
    out = []
    for job in p.jobs:
        start = end_year - job.years
        out.append((f"{job.title} | {job.company} | {start} - {end}", job.bullets))
        end_year = start
        end = start
    return out


def build_lines(p: Person) -> list[tuple[str, str]]:
    """The CV as (kind, text) lines: kind is 'name' | 'heading' | 'text' | 'bullet'."""
    h = HEADINGS[p.lang]
    lines = [
        ("name", p.name),
        ("text", f"{p.email} | {p.phone}"),
        ("text", f"{h['location']}: {p.location}"),
        ("heading", h["summary"]),
        ("text", p.summary),
        ("heading", h["experience"]),
    ]
    for heading, bullets in _job_lines(p):
        lines.append(("text", heading))
        lines += [("bullet", b) for b in bullets]
    if not p.skills_as_table:
        lines += [("heading", h["skills"]), ("text", ", ".join(p.skills))]
    lines.append(("heading", h["education"]))
    lines += [("text", e) for e in p.education]
    if p.projects:
        lines.append(("heading", h["projects"]))
        lines += [("bullet", pr) for pr in p.projects]
    lines += [("heading", h["languages"]), ("text", ", ".join(p.languages))]
    return lines


def _set_rtl(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)


def write_docx(p: Person, path: Path) -> None:
    document = docx.Document()
    rtl = p.lang == "ar"
    for kind, text in build_lines(p):
        if kind == "name":
            para = document.add_paragraph()
            run = para.add_run(text)
            run.bold = True
            run.font.size = Pt(18)
        elif kind == "heading":
            para = document.add_paragraph()
            para.add_run(text).bold = True
        elif kind == "bullet":
            para = document.add_paragraph(f"• {text}")
        else:
            para = document.add_paragraph(text)
        if rtl:
            _set_rtl(para)
    if p.skills_as_table:
        # A two-column table: exercises the parser's table support.
        document.add_paragraph().add_run(HEADINGS[p.lang]["skills"]).bold = True
        rows = [p.skills[i : i + 2] for i in range(0, len(p.skills), 2)]
        table = document.add_table(rows=len(rows), cols=2)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    path.write_bytes(buffer.getvalue())


def write_pdf(p: Person, path: Path) -> None:
    if p.lang != "en":
        raise ValueError("PDF output uses a Latin base-14 font; use DOCX for Arabic CVs")
    doc = pymupdf.open()
    page = doc.new_page()
    y = 60.0
    for kind, text in build_lines(p):
        size = {"name": 18, "heading": 13}.get(kind, 10.5)
        font = "hebo" if kind in ("name", "heading") else "helv"
        if kind == "heading":
            y += 8
        prefix = "- " if kind == "bullet" else ""
        for line in textwrap.wrap(prefix + text, width=95) or [""]:
            if y > 800:
                page = doc.new_page()
                y = 60.0
            page.insert_text((60, y), line, fontsize=size, fontname=font)
            y += size + 5
    path.write_bytes(doc.tobytes())
    doc.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=ROOT / "sample_cvs")
    parser.add_argument(
        "--dataset", type=Path, default=ROOT / "eval" / "datasets" / "sample_queries.json"
    )
    args = parser.parse_args()

    args.out.mkdir(parents=True, exist_ok=True)
    for p in PEOPLE:
        path = args.out / f"{p.cid}.{p.fmt}"
        (write_pdf if p.fmt == "pdf" else write_docx)(p, path)
        print(f"wrote {path.name:<12} {p.lang}  {p.name}  ({p.total_years} yrs)")

    args.dataset.parent.mkdir(parents=True, exist_ok=True)
    args.dataset.write_text(
        json.dumps({"workspace_id": "eval", "queries": QUERIES}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"wrote {args.dataset} ({len(QUERIES)} queries)")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")  # Arabic names on Windows consoles
    main()
